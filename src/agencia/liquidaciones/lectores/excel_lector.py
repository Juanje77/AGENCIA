"""Lector de liquidaciones en Excel (.xlsx)."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from .base import ErrorDeLectura, TablaCruda


def leer_excel(ruta: str | Path, perfil: dict[str, Any] | None = None) -> TablaCruda:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depende del entorno
        raise ErrorDeLectura(
            "Para leer archivos .xlsx hace falta openpyxl. Instalalo con: "
            "pip install openpyxl  (o exporta la liquidacion a CSV)."
        ) from exc

    perfil = perfil or {}
    ruta = Path(ruta)
    avisos: list[str] = []

    try:
        libro = load_workbook(ruta, data_only=True, read_only=True)
    except Exception as exc:
        raise ErrorDeLectura(f"No se pudo abrir {ruta.name}: {exc}") from exc

    try:
        nombre_hoja = perfil.get("hoja")
        if nombre_hoja and nombre_hoja in libro.sheetnames:
            hoja = libro[nombre_hoja]
        else:
            if nombre_hoja:
                avisos.append(
                    f"La hoja {nombre_hoja!r} no existe; se usa {libro.sheetnames[0]!r}."
                )
            hoja = libro[libro.sheetnames[0]]

        filas = [
            list(fila)
            for fila in hoja.iter_rows(values_only=True)
            if any(c is not None and str(c).strip() for c in fila)
        ]
    finally:
        libro.close()

    if not filas:
        raise ErrorDeLectura(f"La hoja {hoja.title!r} de {ruta.name} esta vacia.")

    if len(libro.sheetnames) > 1:
        avisos.append(
            f"El archivo tiene {len(libro.sheetnames)} hojas: {', '.join(libro.sheetnames)}. "
            f"Se leyo {hoja.title!r}."
        )

    return TablaCruda(
        filas=filas, origen=str(ruta), formato="excel", hoja=hoja.title, avisos=avisos
    )
